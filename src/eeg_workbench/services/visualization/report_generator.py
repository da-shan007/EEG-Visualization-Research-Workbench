"""报告生成服务：PDF/HTML/Word 报告生成"""
from __future__ import annotations
from dataclasses import dataclass, field
from typing import Any, Optional, Literal
import time
import json
from pathlib import Path
from datetime import datetime
from io import BytesIO
import base64
import tempfile

from eeg_workbench.models.visualization import (
    ReportConfig, ExportFormat, PlotConfig
)
from eeg_workbench.models.dataset import EEGDataset
from eeg_workbench.services.visualization.plotting import PlottingService
from eeg_workbench.core.events import get_event_bus, EventType, PreprocessingPayload


def _default_report_path(ext: str) -> str:
    """默认输出位置: <cwd>/outputs/reports/eeg_report_<时间戳>.<ext>

    原实现用裸相对路径，报告会直接落在进程 CWD（开发/测试时即项目根目录），
    导致根目录堆积几十份 PDF/HTML。现在统一归档到 outputs/reports/。
    """
    out_dir = Path.cwd() / "outputs" / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    return str(out_dir / f"eeg_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}")


@dataclass
class ReportSection:
    """报告章节"""
    title: str
    content: str
    figures: list = field(default_factory=list)
    tables: list = field(default_factory=list)
    subsections: list = field(default_factory=list)


@dataclass
class ReportResult:
    """报告生成结果"""
    output_path: str
    format: ExportFormat
    processing_time_ms: float
    sections_count: int
    figures_count: int
    tables_count: int
    file_size_bytes: int


class ReportGenerator:
    """报告生成器"""

    def __init__(self):
        self._template_cache: dict = {}

    def generate(
        self,
        dataset: EEGDataset,
        config: ReportConfig,
        sections_data: dict[str, ReportSection] = None,
        **kwargs
    ):
        """生成报告"""
        start_time = time.perf_counter()

        if config.output_format == ExportFormat.PDF:
            return self._generate_pdf(dataset, config, sections_data, **kwargs)
        elif config.output_format == ExportFormat.HTML:
            return self._generate_html(dataset, config, sections_data, **kwargs)
        elif config.output_format == ExportFormat.DOCX:
            return self._generate_docx(dataset, config, sections_data, **kwargs)
        elif config.output_format == ExportFormat.PPTX:
            return self._generate_pptx(dataset, config, sections_data, **kwargs)
        else:
            raise ValueError(f"不支持的输出格式: {config.output_format}")

    def _generate_pdf(
        self,
        dataset: EEGDataset,
        config: ReportConfig,
        sections_data: dict[str, ReportSection] = None,
        **kwargs
    ):
        """生成 PDF 报告 (使用 reportlab)"""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import (
                SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                PageBreak, Image, KeepTogether
            )
            from reportlab.lib.units import cm, inch
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
        except ImportError:
            # 回退到 matplotlib PDF
            return self._generate_pdf_fallback(config)

        output_path = config.output_path or _default_report_path("pdf")
        
        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            rightMargin=2*cm,
            leftMargin=2*cm,
            topMargin=2*cm,
            bottomMargin=2*cm
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle', parent=styles['Title'],
            fontSize=18, spaceAfter=20, alignment=TA_CENTER
        )
        heading_style = ParagraphStyle(
            'CustomHeading', parent=styles['Heading1'],
            fontSize=14, spaceAfter=10, textColor=colors.darkblue
        )
        body_style = ParagraphStyle(
            'CustomBody', parent=styles['Normal'],
            fontSize=11, leading=14, alignment=TA_JUSTIFY
        )

        story = []

        # 标题页
        story.append(Paragraph(config.title, title_style))
        story.append(Spacer(1, 20))
        story.append(Paragraph(f"作者: {config.author or '未知'}", styles['Normal']))
        story.append(Paragraph(f"机构: {config.institution or '未知'}", styles['Normal']))
        story.append(Paragraph(f"日期: {datetime.now().strftime(config.date_format)}", styles['Normal']))
        story.append(PageBreak())

        # 目录
        if config.include_toc:
            story.append(Paragraph("目录", styles['Heading1']))
            story.append(Spacer(1, 12))
            for i, section_name in enumerate(config.sections, 1):
                story.append(Paragraph(f"{i}. {section_name}", styles['Normal']))
            story.append(PageBreak())

        # 章节内容
        if sections_data:
            for section_name in config.sections:
                if section_name in sections_data:
                    section = sections_data[section_name]
                    story.append(Paragraph(section.title, heading_style))
                    story.append(Spacer(1, 6))
                    
                    # 章节内容
                    for para in section.content.split('\n'):
                        if para.strip():
                            story.append(Paragraph(para, body_style))
                            story.append(Spacer(1, 6))

                    # 图表
                    for fig_result in section.figures:
                        if fig_result.figure:
                            img_path = self._save_figure_temp(fig_result.figure)
                            if img_path:
                                img = Image(img_path, width=config.figure_width_cm*cm, height=config.figure_height_cm*cm)
                                story.append(Spacer(1, 12))
                                story.append(img)
                                story.append(Spacer(1, 12))

                    # 表格
                    for table_data in section.tables:
                        if table_data:
                            table = self._create_table(table_data, config)
                            story.append(Spacer(1, 12))
                            story.append(table)
                            story.append(Spacer(1, 12))

                    story.append(PageBreak())

        doc.build(story)

        file_size = Path(output_path).stat().st_size
        
        return ReportResult(
            output_path=output_path,
            format=ExportFormat.PDF,
            processing_time_ms=0,
            sections_count=len(config.sections),
            figures_count=sum(len(s.figures) for s in sections_data.values()) if sections_data else 0,
            tables_count=sum(len(s.tables) for s in sections_data.values()) if sections_data else 0,
            file_size_bytes=file_size
        )

    def _generate_pdf_fallback(self, config):
        """使用 matplotlib 回退生成 PDF"""
        from matplotlib.backends.backend_pdf import PdfPages
        import matplotlib.pyplot as plt

        output_path = config.output_path or _default_report_path("pdf")
        
        with PdfPages(output_path) as pdf:
            # 标题页
            fig, ax = plt.subplots(figsize=(8.27, 11.69))  # A4
            ax.axis('off')
            ax.text(0.5, 0.7, "EEG 分析报告", ha='center', va='center', fontsize=24, fontweight='bold')
            ax.text(0.5, 0.6, f"作者: {config.author or '未知'}", ha='center', fontsize=14)
            ax.text(0.5, 0.55, f"机构: {config.institution or '未知'}", ha='center', fontsize=12)
            ax.text(0.5, 0.5, f"日期: {datetime.now().strftime('%Y-%m-%d')}", ha='center', fontsize=12)
            pdf.savefig(fig)
            plt.close(fig)

            # 章节页
            for section_name in ["overview", "methods", "preprocessing", "erp", "time_frequency", 
                                "connectivity", "source_localization", "statistics"]:
                fig, ax = plt.subplots(figsize=(8.27, 11.69))
                ax.axis('off')
                ax.text(0.5, 0.9, section_name.replace('_', ' ').title(), ha='center', fontsize=18, fontweight='bold')
                ax.text(0.5, 0.5, f"{section_name} 内容待填充...", ha='center', fontsize=12)
                pdf.savefig(fig)
                plt.close(fig)

        file_size = Path(output_path).stat().st_size
        
        return ReportResult(
            output_path=output_path,
            format=ExportFormat.PDF,
            processing_time_ms=0,
            sections_count=8,
            figures_count=0,
            tables_count=0,
            file_size_bytes=file_size
        )

    def _generate_html(
        self,
        dataset: EEGDataset,
        config: ReportConfig,
        sections_data: dict[str, ReportSection] = None,
        **kwargs
    ):
        """生成 HTML 报告 (交互式)"""
        output_path = config.output_path or _default_report_path("html")

        html_template = self._get_html_template()
        
        # 渲染各章节
        sections_html = []
        if sections_data:
            for section_name in config.sections:
                if section_name in sections_data:
                    section = sections_data[section_name]
                    sec_html = f"<section><h2>{section.title}</h2>"
                    
                    # 文本内容
                    for para in section.content.split('\n'):
                        if para.strip():
                            sec_html += f"<p>{para}</p>"
                    
                    # 图表
                    for fig_result in section.figures:
                        if fig_result.figure:
                            img_b64 = self._fig_to_base64(fig_result.figure)
                            sec_html += f'<div class="figure"><img src="data:image/png;base64,{img_b64}" style="max-width:100%;"></div>'
                    
                    # 表格
                    for table_data in section.tables:
                        sec_html += self._table_to_html(table_data)
                    
                    sec_html += "</section>"
                    sections_html.append(sec_html)

        html_content = self._get_html_template().format(
            title=config.title,
            author=config.author,
            institution=config.institution,
            date=datetime.now().strftime(config.date_format),
            sections=''.join(sections_html),
            custom_css=config.custom_css
        )

        output_path = config.output_path or _default_report_path("html")
        Path(output_path).write_text(html_content, encoding='utf-8')

        file_size = Path(output_path).stat().st_size

        return ReportResult(
            output_path=output_path,
            format=ExportFormat.HTML,
            processing_time_ms=0,
            sections_count=len(config.sections),
            figures_count=0,
            tables_count=0,
            file_size_bytes=file_size
        )

    def _get_html_template(self) -> str:
        """获取 HTML 模板"""
        template = """
<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <style>
        body {{ font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; line-height: 1.6; max-width: 900px; margin: 0 auto; padding: 20px; }}
        h1 {{ color: #2c3e50; border-bottom: 2px solid #3498db; padding-bottom: 10px; }}
        h2 {{ color: #34495e; margin-top: 30px; }}
        p {{ text-align: justify; }}
        .figure {{ text-align: center; margin: 20px 0; }}
        .figure img {{ max-width: 100%; height: auto; box-shadow: 0 2px 8px rgba(0,0,0,0.1); }}
        table {{ border-collapse: collapse; width: 100%; margin: 20px 0; }}
        th, td {{ border: 1px solid #ddd; padding: 12px; text-align: left; }}
        th {{ background-color: #3498db; color: white; }}
        tr:nth-child(even) {{ background-color: #f2f2f2; }}
        .toc {{ background: #f8f9fa; padding: 20px; border-radius: 8px; margin: 20px 0; }}
        .toc ul {{ list-style: none; padding-left: 20px; }}
        .toc a {{ text-decoration: none; color: #3498db; }}
        {custom_css}
    </style>
    <script src="https://cdn.plot.ly/plotly-latest.min.js"></script>
</head>
<body>
    <h1>{title}</h1>
    <div class="meta">作者: {author} | 机构: {institution} | 日期: {date}</div>
    
    <div class="toc">
        <h2>目录</h2>
        <ul>
            <li><a href="#overview">概览</a></li>
            <li><a href="#methods">方法</a></li>
            <li><a href="#preprocessing">预处理</a></li>
            <li><a href="#erp">ERP 分析</a></li>
            <li><a href="#time_frequency">时频分析</a></li>
            <li><a href="#connectivity">连通性分析</a></li>
            <li><a href="#source_localization">源定位</a></li>
            <li><a href="#statistics">统计分析</a></li>
            <li><a href="#conclusion">结论</a></li>
        </ul>
    </div>

    {sections}

    <hr>
    <footer style="text-align: center; color: #7f8c8d; margin-top: 40px;">
        <p>由 EEG Visualization Research Workbench 生成</p>
    </footer>
</body>
</html>
        """
        return template

    def _fig_to_base64(self, fig) -> str:
        """将 matplotlib 图形转为 base64"""
        buf = BytesIO()
        fig.savefig(buf, format='png', dpi=150, bbox_inches='tight')
        buf.seek(0)
        return base64.b64encode(buf.read()).decode('utf-8')

    def _table_to_html(self, table_data: dict) -> str:
        """将表格数据转为 HTML"""
        if not table_data:
            return ""
        
        html = '<table>'
        if 'headers' in table_data and table_data['headers']:
            html += '<tr>' + ''.join(f'<th>{h}</th>' for h in table_data['headers']) + '</tr>'
        if 'rows' in table_data:
            for row in table_data['rows']:
                html += '<tr>' + ''.join(f'<td>{cell}</td>' for cell in row) + '</tr>'
        html += '</table>'
        return html

    def _create_table(self, table_data: dict, config):
        """创建 reportlab 表格"""
        from reportlab.platypus import Table, TableStyle
        from reportlab.lib import colors
        
        data = [table_data.get('headers', [])] + table_data.get('rows', [])
        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f2f2f2')),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        return table

    def _save_figure_temp(self, fig) -> str:
        """保存图形为临时文件"""
        tmp = tempfile.NamedTemporaryFile(suffix='.png', delete=False)
        fig.savefig(tmp.name, dpi=150, bbox_inches='tight')
        return tmp.name

    def _generate_docx(self, dataset, config, sections_data, **kwargs):
        """生成 DOCX 报告 (需要 python-docx: pip install python-docx)"""
        try:
            from docx import Document
            from docx.shared import Inches, Pt
        except ImportError:
            raise RuntimeError("DOCX 报告需要 python-docx: pip install python-docx")
        start_time = time.perf_counter()
        output_path = config.output_path or _default_report_path("docx")

        doc = Document()
        doc.add_heading(config.title, level=0)
        doc.add_paragraph(f"作者: {config.author} | 机构: {config.institution} | 日期: {datetime.now().strftime(config.date_format)}")

        n_figs = n_tables = 0
        if sections_data:
            for section_name in config.sections:
                if section_name not in sections_data:
                    continue
                section = sections_data[section_name]
                doc.add_heading(section.title, level=1)
                for para in section.content.split('\n'):
                    if para.strip():
                        doc.add_paragraph(para)
                for fig_result in section.figures:
                    fig = getattr(fig_result, 'figure', None)
                    if fig is not None:
                        tmp = self._save_figure_temp(fig)
                        try:
                            doc.add_picture(tmp, width=Inches(6))
                            n_figs += 1
                        finally:
                            Path(tmp).unlink(missing_ok=True)
                for table_data in section.tables:
                    headers = table_data.get('headers', []) if table_data else []
                    rows = table_data.get('rows', []) if table_data else []
                    if not rows:
                        continue
                    table = doc.add_table(rows=1 + len(rows), cols=max(len(headers), len(rows[0])))
                    table.style = 'Light Grid'
                    for j, h in enumerate(headers):
                        table.cell(0, j).text = str(h)
                    for i, row in enumerate(rows):
                        for j, cell in enumerate(row):
                            table.cell(i + 1, j).text = str(cell)
                    n_tables += 1
        doc.save(output_path)
        return ReportResult(
            output_path=output_path, format=ExportFormat.DOCX,
            processing_time_ms=(time.perf_counter() - start_time) * 1000,
            sections_count=len(config.sections),
            figures_count=n_figs, tables_count=n_tables,
            file_size_bytes=Path(output_path).stat().st_size,
        )

    def _generate_pptx(self, dataset, config, sections_data, **kwargs):
        """生成 PPTX 报告 (需要 python-pptx: pip install python-pptx)"""
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError:
            raise RuntimeError("PPTX 报告需要 python-pptx: pip install python-pptx")
        start_time = time.perf_counter()
        output_path = config.output_path or _default_report_path("pptx")

        prs = Presentation()
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = config.title
        title_slide.placeholders[1].text = f"{config.author} | {config.institution}"

        n_figs = 0
        if sections_data:
            for section_name in config.sections:
                if section_name not in sections_data:
                    continue
                section = sections_data[section_name]
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = section.title
                body = slide.placeholders[1].text_frame
                body.clear()
                for para in section.content.split('\n'):
                    if para.strip():
                        body.add_paragraph().text = para
                for fig_result in section.figures:
                    fig = getattr(fig_result, 'figure', None)
                    if fig is not None:
                        tmp = self._save_figure_temp(fig)
                        try:
                            slide.shapes.add_picture(tmp, Inches(0.5), Inches(2.5), width=Inches(9))
                            n_figs += 1
                        finally:
                            Path(tmp).unlink(missing_ok=True)
                        break
        prs.save(output_path)
        return ReportResult(
            output_path=output_path, format=ExportFormat.PPTX,
            processing_time_ms=(time.perf_counter() - start_time) * 1000,
            sections_count=len(config.sections),
            figures_count=n_figs, tables_count=0,
            file_size_bytes=Path(output_path).stat().st_size,
        )


def generate_report(
    dataset: EEGDataset,
    config: ReportConfig,
    sections_data: dict[str, ReportSection] = None,
    **kwargs
) -> ReportResult:
    """生成报告的便捷函数"""
    generator = ReportGenerator()
    return generator.generate(dataset, config, sections_data, **kwargs)


def generate_html_report(
    dataset: EEGDataset,
    config: ReportConfig,
    sections_data: dict[str, ReportSection] = None,
    **kwargs
) -> ReportResult:
    """生成 HTML 报告的便捷函数"""
    config.output_format = ExportFormat.HTML
    generator = ReportGenerator()
    return generator.generate(dataset, config, sections_data, **kwargs)